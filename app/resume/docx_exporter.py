from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

from app.models.candidate import CandidateProfile
from app.models.job import JobPosting
from app.models.tailored_resume import TailoredResume
from app.resume.project_bullets import project_bullets


NAVY = RGBColor(0x1F, 0x2A, 0x44)
MUTED = RGBColor(0x5B, 0x67, 0x7A)
RULE = RGBColor(0xC5, 0xCD, 0xD8)


class ResumeExporter:

    def __init__(
        self,
        output_dir: str = "data/resumes",
    ):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(
            parents=True,
            exist_ok=True,
        )

    def export(
        self,
        candidate: CandidateProfile,
        job: JobPosting,
        resume: TailoredResume,
        dest: Path | None = None,
    ) -> str:

        document = Document()
        self._set_page(document)
        self._set_base_style(document)

        self._add_header(document, candidate, job)
        self._add_section(document, "Professional Summary")
        self._add_body(document, resume.summary)

        self._add_section(document, "Technical Skills")
        self._add_skills(
            document,
            candidate.skills or resume.skills,
        )

        domains = candidate.domains or resume.competencies

        if domains:
            self._add_section(
                document,
                "Core Competencies",
            )
            for domain in domains:
                name = str(domain or "").strip()

                if name:
                    self._add_bullet(document, name)

        self._add_section(document, "Work Experience")

        if resume.experiences:
            for block in resume.experiences:
                self._add_job_block(document, block)
        else:
            for highlight in resume.experience_highlights:
                self._add_bullet(document, highlight)

        if resume.project_highlights:
            self._add_section(document, "Projects")

            for highlight in resume.project_highlights:
                self._add_project(document, highlight)

        if candidate.education:
            self._add_section(document, "Education")

            for item in candidate.education:
                education = document.add_paragraph()
                education.paragraph_format.space_after = (
                    Pt(6)
                )
                degree = education.add_run(
                    item.degree
                )
                degree.bold = True
                degree.font.size = Pt(11)
                degree.font.color.rgb = NAVY

                school = document.add_paragraph()
                school.paragraph_format.space_after = (
                    Pt(12)
                )
                school_run = school.add_run(
                    item.institution
                )
                school_run.font.size = Pt(10.5)
                school_run.font.color.rgb = MUTED

                if item.year:
                    school_run2 = school.add_run(
                        f"  ·  {item.year}"
                    )
                    school_run2.font.size = Pt(10.5)
                    school_run2.font.color.rgb = MUTED

        if candidate.certifications:
            self._add_section(document, "Certifications")

            for item in candidate.certifications:
                cert = re.sub(
                    r"^[\s•\-\u2022*]+",
                    "",
                    str(item or ""),
                ).strip()

                if cert:
                    self._add_skill_line(document, cert)

        path = dest or (
            self.output_dir
            / self.filename(job)
        )
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        document.save(path)

        return str(path)

    @staticmethod
    def filename(job: JobPosting) -> str:

        safe_title = "".join(
            character
            if character.isalnum()
            else "_"
            for character in job.title
        )[:40]

        return f"{job.job_id}_{safe_title}.docx"

    @staticmethod
    def _set_page(document: Document) -> None:

        for section in document.sections:
            section.top_margin = Inches(0.85)
            section.bottom_margin = Inches(0.85)
            section.left_margin = Inches(0.9)
            section.right_margin = Inches(0.9)

    @staticmethod
    def _set_base_style(document: Document) -> None:

        style = document.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.font.color.rgb = NAVY
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.15

    def _add_header(
        self,
        document: Document,
        candidate: CandidateProfile,
        job: JobPosting,
    ) -> None:

        name = document.add_paragraph()
        name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name.paragraph_format.space_after = Pt(8)
        run = name.add_run(candidate.name)
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"

        self._add_contact_line(document, candidate)

        title = document.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.space_after = Pt(14)
        title_run = title.add_run(job.title)
        title_run.font.size = Pt(12)
        title_run.italic = True
        title_run.font.color.rgb = MUTED

        self._add_rule(document)

    def _add_contact_line(
        self,
        document: Document,
        candidate: CandidateProfile,
    ) -> None:

        items = []

        if candidate.email:
            items.append(("text", candidate.email, ""))

        if candidate.phone:
            items.append(("text", candidate.phone, ""))

        if candidate.preferred_locations:
            items.append(
                ("text", candidate.preferred_locations[0], "")
            )

        if candidate.linkedin:
            items.append(
                (
                    "link",
                    "LinkedIn",
                    self._absolute_url(candidate.linkedin),
                )
            )

        if candidate.github:
            items.append(
                (
                    "link",
                    "GitHub",
                    self._absolute_url(candidate.github),
                )
            )

        if not items:
            return

        contact = document.add_paragraph()
        contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact.paragraph_format.space_after = Pt(8)

        for index, (kind, label, url) in enumerate(items):
            if index:
                sep = contact.add_run("  ·  ")
                sep.font.size = Pt(10)
                sep.font.color.rgb = MUTED

            if kind == "link":
                self._add_hyperlink(contact, label, url)
            else:
                run = contact.add_run(label)
                run.font.size = Pt(10)
                run.font.color.rgb = MUTED

    @staticmethod
    def _absolute_url(value: str) -> str:

        text = str(value or "").strip()

        if text.startswith("http://") or text.startswith("https://"):
            return text

        return f"https://{text.lstrip('/')}"

    @staticmethod
    def _add_hyperlink(
        paragraph,
        text: str,
        url: str,
    ) -> None:

        r_id = paragraph.part.relate_to(
            url,
            RT.HYPERLINK,
            is_external=True,
        )
        label = (
            str(text)
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
        hyperlink = parse_xml(
            f'<w:hyperlink {nsdecls("w", "r")} r:id="{r_id}" '
            f'w:history="1">'
            f"<w:r>"
            f"<w:rPr>"
            f'<w:rStyle w:val="Hyperlink"/>'
            f'<w:color w:val="0563C1"/>'
            f'<w:u w:val="single"/>'
            f'<w:sz w:val="20"/>'
            f"</w:rPr>"
            f'<w:t xml:space="preserve">{label}</w:t>'
            f"</w:r>"
            f"</w:hyperlink>"
        )
        paragraph._p.append(hyperlink)

    def _add_section(
        self,
        document: Document,
        heading: str,
    ) -> None:

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = (
            Pt(18)
        )
        paragraph.paragraph_format.space_after = (
            Pt(10)
        )
        run = paragraph.add_run(heading.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"
        self._underline(paragraph)

    def _add_skills(
        self,
        document: Document,
        skills: list[str],
    ) -> None:

        if not skills:
            return

        for skill in skills:
            self._add_skill_line(document, skill)

    def _add_skill_line(
        self,
        document: Document,
        text: str,
    ) -> None:

        line = str(text or "").strip()

        if not line:
            return

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.15
        paragraph.paragraph_format.left_indent = Inches(0)

        label, separator, rest = line.partition(":")

        if separator and rest.strip():
            head = paragraph.add_run(f"{label.strip()}:")
            head.bold = True
            head.font.size = Pt(11)
            head.font.color.rgb = NAVY
            head.font.name = "Calibri"
            body = paragraph.add_run(f" {rest.strip()}")
            body.font.size = Pt(11)
            body.font.color.rgb = NAVY
            body.font.name = "Calibri"
            return

        run = paragraph.add_run(line)
        run.font.size = Pt(11)
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"

    def _add_job_block(
        self,
        document: Document,
        block,
    ) -> None:

        header = document.add_paragraph()
        header.paragraph_format.space_before = Pt(16)
        header.paragraph_format.space_after = Pt(4)
        self._right_tab(header)

        company = header.add_run(block.company)
        company.bold = True
        company.font.size = Pt(12)
        company.font.color.rgb = NAVY

        if block.dates:
            header.add_run("\t")
            dates = header.add_run(block.dates)
            dates.font.size = Pt(10)
            dates.font.color.rgb = MUTED

        role = document.add_paragraph()
        role.paragraph_format.space_after = Pt(10)
        role_run = role.add_run(block.role)
        role_run.italic = True
        role_run.font.size = Pt(10.5)
        role_run.font.color.rgb = MUTED

        for bullet in block.bullets:
            self._add_bullet(document, bullet)

    def _add_project(
        self,
        document: Document,
        highlight: str,
    ) -> None:

        name, bullets = self._project_parts(highlight)
        tools = []
        body = []

        for bullet in bullets:
            if re.match(r"^(stack|tools)\b", bullet, re.I):
                tools.append(bullet)
            else:
                body.append(bullet)

        if not body and not tools:
            self._add_bullet(document, name)
            return

        header = document.add_paragraph()
        header.paragraph_format.space_before = Pt(6)
        header.paragraph_format.space_after = Pt(2)
        run = header.add_run(name)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = NAVY

        if tools:
            line = document.add_paragraph()
            line.paragraph_format.space_after = Pt(8)
            tools_run = line.add_run(tools[0])
            tools_run.italic = True
            tools_run.font.size = Pt(10.5)
            tools_run.font.color.rgb = MUTED

        for bullet in body:
            self._add_bullet(document, bullet)

    @staticmethod
    def _project_parts(
        highlight: str,
    ) -> tuple[str, list[str]]:

        text = str(highlight or "").strip()

        if not text:
            return "", []

        if "\n" in text:
            lines = [
                line.strip()
                for line in text.splitlines()
                if line.strip()
            ]
            return lines[0], lines[1:]

        name, separator, description = text.partition(": ")

        if not separator:
            return text, []

        return name, project_bullets(description)

    def _add_body(
        self,
        document: Document,
        text: str,
    ) -> None:

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = (
            Pt(12)
        )
        paragraph.paragraph_format.line_spacing = 1.2
        run = paragraph.add_run(text)
        run.font.size = Pt(11)
        run.font.color.rgb = NAVY

    def _add_bullet(
        self,
        document: Document,
        text: str,
    ) -> None:

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = (
            Pt(3)
        )
        paragraph.paragraph_format.space_after = (
            Pt(8)
        )
        paragraph.paragraph_format.line_spacing = (
            1.2
        )
        paragraph.paragraph_format.left_indent = (
            Inches(0.2)
        )
        run = paragraph.add_run(f"• {text}")
        run.font.size = Pt(11)
        run.font.color.rgb = NAVY
        run.font.name = "Calibri"

    @staticmethod
    def _add_rule(document: Document) -> None:

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = (
            Pt(4)
        )
        pPr = paragraph._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "C5CDD8")
        pBdr.append(bottom)
        pPr.append(pBdr)

    @staticmethod
    def _underline(paragraph) -> None:

        pPr = paragraph._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "1F2A44")
        pBdr.append(bottom)
        pPr.append(pBdr)

    @staticmethod
    def _right_tab(paragraph) -> None:

        pPr = paragraph._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "right")
        tab.set(qn("w:pos"), "9360")
        tabs.append(tab)
        pPr.append(tabs)

